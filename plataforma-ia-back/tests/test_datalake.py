"""Tests del lakehouse-lite: ingest + query + export.

Usan storage_service en modo disco local (config por defecto sin conexión
Azure); cada test genera un conversation_id único para evitar colisiones.
"""
from __future__ import annotations

import uuid
from io import BytesIO

import pytest
import xlsxwriter

from app.agents.datalake import export_dataset, ingest_file, query_dataset
from app.agents.datalake.catalog import DatasetCatalog
from app.services.storage_service import storage_service


def _fresh_conv() -> str:
    return f"test{uuid.uuid4().hex[:10]}"


def _make_xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    buf = BytesIO()
    # strings_to_formulas=False evita que xlsxwriter interprete '=CMD|...' como fórmula
    wb = xlsxwriter.Workbook(
        buf,
        {"in_memory": True, "strings_to_formulas": False},
    )
    for sheet_name, rows in sheets.items():
        ws = wb.add_worksheet(sheet_name)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if isinstance(value, str):
                    ws.write_string(r, c, value)
                else:
                    ws.write(r, c, value)
    wb.close()
    return buf.getvalue()


def _cleanup(conv_id: str, datasets: list[str]) -> None:
    for ds_id in datasets:
        DatasetCatalog.delete(conv_id, ds_id)


def test_ingest_csv_roundtrip():
    conv = _fresh_conv()
    csv_bytes = b"colA,colB\n1,foo\n2,bar\n3,baz\n"
    result = ingest_file(csv_bytes, "data.csv", conv)
    assert result["success"], result
    assert len(result["datasets"]) == 1
    meta = result["datasets"][0]
    assert meta["rows"] == 3
    assert set(meta["schema"].keys()) == {"colA", "colB"}

    qres = query_dataset(conv, meta["dataset_id"], "SELECT COUNT(*) AS n FROM dataset")
    assert qres["success"]
    assert qres["rows"][0]["n"] == 3

    _cleanup(conv, [meta["dataset_id"]])


def test_ingest_xlsx_multisheet():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({
        "Ventas": [["division", "valor"], ["Norte", 100], ["Sur", 200]],
        "Gastos": [["div", "total"], ["Norte", 50]],
        "Vacia": [["h"]],
    })
    result = ingest_file(content, "multi.xlsx", conv)
    assert result["success"], result
    # 'Vacia' se omite (sin filas de datos)
    assert len(result["datasets"]) == 2
    names = sorted(d["sheet"] for d in result["datasets"])
    assert names == ["Gastos", "Ventas"]

    _cleanup(conv, [d["dataset_id"] for d in result["datasets"]])


def test_ingest_tsv_gzip():
    import gzip
    conv = _fresh_conv()
    tsv = "a\tb\n1\tx\n2\ty\n"
    gz_bytes = gzip.compress(tsv.encode("utf-8"))
    result = ingest_file(gz_bytes, "data.tsv.gz", conv)
    assert result["success"], result
    assert result["datasets"][0]["rows"] == 2
    _cleanup(conv, [result["datasets"][0]["dataset_id"]])


def test_ingest_zip_bomb_rejected(monkeypatch):
    # Forzamos un ratio > 200 con un stub de zip_bomb_guard indirecto:
    # construimos un xlsx minúsculo y parchamos el guard para simular bomb.
    from app.agents import file_security

    def fake_guard(*args, **kwargs):
        raise file_security.FileSecurityError("ratio 9999:1 excede 200:1")

    monkeypatch.setattr("app.agents.datalake.ingest.zip_bomb_guard", fake_guard)

    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [["a"], [1]]})
    result = ingest_file(content, "bomb.xlsx", conv)
    assert not result["success"]
    assert "bloqueado" in result["error"]


def test_query_blocks_attach():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [["a", "b"], [1, 2]]})
    result = ingest_file(content, "s.xlsx", conv)
    ds = result["datasets"][0]["dataset_id"]
    try:
        for bad in [
            "ATTACH '/etc/passwd' AS x",
            "PRAGMA database_list",
            "COPY dataset TO 'out.csv'",
            "DROP TABLE dataset",
            "UPDATE dataset SET a=0",
            "INSTALL httpfs",
            "LOAD 'x.duckdb_extension'",
        ]:
            res = query_dataset(conv, ds, bad)
            assert not res["success"], f"debió rechazar: {bad}"
            assert "bloqueadas" in res["error"] or "vacía" in res["error"]
    finally:
        _cleanup(conv, [ds])


def test_query_blocks_multiple_statements():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [["a"], [1]]})
    ds = ingest_file(content, "m.xlsx", conv)["datasets"][0]["dataset_id"]
    try:
        res = query_dataset(conv, ds, "SELECT 1; SELECT 2")
        assert not res["success"]
        assert "múltiples" in res["error"]
    finally:
        _cleanup(conv, [ds])


def test_query_limit_truncated():
    conv = _fresh_conv()
    header = [["idx", "val"]]
    rows = [[i, i * 2] for i in range(1_200)]
    content = _make_xlsx_bytes({"S": header + rows})
    ds = ingest_file(content, "big.xlsx", conv)["datasets"][0]["dataset_id"]
    try:
        res = query_dataset(conv, ds, "SELECT * FROM dataset", limit=500)
        assert res["success"]
        assert res["row_count"] == 500
        assert res["truncated"] is True
    finally:
        _cleanup(conv, [ds])


def test_query_aggregation():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [
        ["div", "v"], ["N", 10], ["S", 20], ["N", 30], ["S", 5],
    ]})
    ds = ingest_file(content, "agg.xlsx", conv)["datasets"][0]["dataset_id"]
    try:
        res = query_dataset(conv, ds,
            "SELECT div, SUM(v) AS t FROM dataset GROUP BY 1 ORDER BY 2 DESC")
        assert res["success"]
        assert res["rows"] == [{"div": "N", "t": 40}, {"div": "S", "t": 25}]
    finally:
        _cleanup(conv, [ds])


def test_export_xlsx_has_branding():
    from openpyxl import load_workbook

    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [["col", "valor"], ["A", 1], ["B", 2]]})
    ds = ingest_file(content, "e.xlsx", conv)["datasets"][0]["dataset_id"]
    filename = f"exp_{uuid.uuid4().hex[:6]}.xlsx"
    try:
        res = export_dataset(conv, ds, "xlsx", {"filename": filename})
        assert res["success"]
        assert res["filename"] == filename

        downloaded = storage_service.download_blob(res["filepath"])
        wb = load_workbook(BytesIO(downloaded))
        props = wb.properties
        assert props.creator == "Plataforma IA"
    finally:
        storage_service.delete_blob(f"generated/{filename}")
        _cleanup(conv, [ds])


def test_export_parquet_passthrough():
    import polars as pl

    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [["c1", "c2"], [1.5, "foo"], [2.5, "bar"]]})
    ds = ingest_file(content, "p.xlsx", conv)["datasets"][0]["dataset_id"]
    filename = f"out_{uuid.uuid4().hex[:6]}.parquet"
    try:
        res = export_dataset(conv, ds, "parquet", {"filename": filename})
        assert res["success"]

        downloaded = storage_service.download_blob(res["filepath"])
        df = pl.read_parquet(BytesIO(downloaded))
        assert df.columns == ["c1", "c2"]
        assert df.height == 2
    finally:
        storage_service.delete_blob(f"generated/{filename}")
        _cleanup(conv, [ds])


def test_export_csv_with_bom_and_sanitization():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [
        ["name", "formula"],
        ["Alice", "=CMD|'/c calc'"],
        ["Bob", "plain"],
    ]})
    ds = ingest_file(content, "inj.xlsx", conv)["datasets"][0]["dataset_id"]
    filename = f"csv_{uuid.uuid4().hex[:6]}.csv"
    try:
        res = export_dataset(conv, ds, "csv", {"filename": filename})
        assert res["success"]

        downloaded = storage_service.download_blob(res["filepath"])
        # BOM UTF-8
        assert downloaded.startswith(b"\xef\xbb\xbf")
        text = downloaded.decode("utf-8-sig")
        # La celda maliciosa debe estar prefijada con apóstrofo
        assert "'=CMD|" in text or "'=CMD" in text
    finally:
        storage_service.delete_blob(f"generated/{filename}")
        _cleanup(conv, [ds])


def test_export_with_filter_sql_and_columns():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [
        ["id", "category", "value"],
        [1, "A", 10], [2, "B", 20], [3, "A", 30], [4, "C", 40],
    ]})
    ds = ingest_file(content, "f.xlsx", conv)["datasets"][0]["dataset_id"]
    filename = f"filt_{uuid.uuid4().hex[:6]}.json"
    try:
        res = export_dataset(conv, ds, "json", {
            "filename": filename,
            "filter_sql": "SELECT id, value FROM dataset WHERE category = 'A'",
            "sort_by": "value",
            "descending": True,
        })
        assert res["success"]
        assert res["rows"] == 2
        assert res["columns"] == ["id", "value"]

        downloaded = storage_service.download_blob(res["filepath"])
        import json
        data = json.loads(downloaded.decode("utf-8"))
        assert [d["id"] for d in data] == [3, 1]
    finally:
        storage_service.delete_blob(f"generated/{filename}")
        _cleanup(conv, [ds])


def test_list_for_conversation():
    conv = _fresh_conv()
    ids: list[str] = []
    try:
        for i in range(3):
            content = _make_xlsx_bytes({f"Hoja{i}": [["a"], [1]]})
            res = ingest_file(content, f"f{i}.xlsx", conv)
            ids.append(res["datasets"][0]["dataset_id"])
        listed = DatasetCatalog.list_for_conversation(conv)
        assert len(listed) == 3
        assert {m.dataset_id for m in listed} == set(ids)
    finally:
        _cleanup(conv, ids)


def test_unsupported_format():
    conv = _fresh_conv()
    res = ingest_file(b"random binary \x00\x01", "weird.xyz", conv)
    assert not res["success"]
    assert "Formato" in res["error"]


def test_export_rejects_unsupported_format():
    conv = _fresh_conv()
    content = _make_xlsx_bytes({"S": [["a"], [1]]})
    ds = ingest_file(content, "s.xlsx", conv)["datasets"][0]["dataset_id"]
    try:
        res = export_dataset(conv, ds, "pdf")
        assert not res["success"]
        assert "no soportado" in res["error"]
    finally:
        _cleanup(conv, [ds])
