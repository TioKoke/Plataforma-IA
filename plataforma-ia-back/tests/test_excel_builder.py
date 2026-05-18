"""Tests para el constructor de reportes Excel Plataforma IA."""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.agents.excel import (
    ChartSpec,
    ColumnSpec,
    ExcelReportSpec,
    KPISpec,
    SheetSpec,
    SummarySheetSpec,
    build_excel_report,
    sanitize_cell,
)
from app.agents.excel.auto_style import (
    infer_chart,
    infer_kpis_from_rows,
    infer_report_type,
)
from app.agents.excel.builder import infer_spec_from_rows
from app.agents.excel.theme import BRAND_NAVY, BRAND_ORANGE


# ── sanitize ────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("value,expected", [
    ("=cmd|'/c calc'!A1", "'=cmd|'/c calc'!A1"),
    ("+HYPERLINK(\"x\",\"x\")", "'+HYPERLINK(\"x\",\"x\")"),
    ("-2+3", "'-2+3"),
    ("@SUM(A1)", "'@SUM(A1)"),
    ("\tmaliciouos", "'\tmaliciouos"),
    ("normal text", "normal text"),
    ("", ""),
])
def test_sanitize_cell_prefixes_dangerous_strings(value, expected):
    assert sanitize_cell(value) == expected


@pytest.mark.parametrize("value", [123, 45.6, True, None])
def test_sanitize_cell_passthrough_non_strings(value):
    assert sanitize_cell(value) is value


# ── build_excel_report ──────────────────────────────────────────────────────

def _minimal_spec() -> ExcelReportSpec:
    return ExcelReportSpec(
        filename="test.xlsx",
        sheets=[SheetSpec(
            name="Datos",
            columns=[
                ColumnSpec(header="División", key="div", type="text"),
                ColumnSpec(header="Toneladas", key="ton", type="integer"),
                ColumnSpec(header="USD MM", key="usd", type="currency"),
            ],
            rows=[
                {"div": "Norte", "ton": 82500, "usd": 612.4},
                {"div": "RT", "ton": 67800, "usd": 501.2},
            ],
            title="Producción Q1 2026",
        )],
        author="Jorge Tapia",
        division="Norte",
    )


def test_build_minimal_report_returns_openable_bytes():
    data = build_excel_report(_minimal_spec())
    assert isinstance(data, bytes)
    assert len(data) > 0

    wb = load_workbook(BytesIO(data))
    assert "Datos" in wb.sheetnames


def test_column_types_produce_expected_number_formats():
    spec = ExcelReportSpec(
        filename="types.xlsx",
        sheets=[SheetSpec(
            name="Tipos",
            columns=[
                ColumnSpec(header="Txt", key="t", type="text"),
                ColumnSpec(header="Int", key="i", type="integer"),
                ColumnSpec(header="Num", key="n", type="number"),
                ColumnSpec(header="Pct", key="p", type="percent"),
                ColumnSpec(header="Cur", key="c", type="currency"),
                ColumnSpec(header="Fch", key="d", type="date"),
            ],
            rows=[{"t": "a", "i": 1, "n": 1.5, "p": 0.2, "c": 100, "d": "2026-04-20"}],
        )],
    )
    wb = load_workbook(BytesIO(build_excel_report(spec)))
    ws = wb["Tipos"]

    data_row = 2
    assert "#,##0" in ws.cell(row=data_row, column=2).number_format
    assert "0.0%" in ws.cell(row=data_row, column=4).number_format
    assert "USD" in ws.cell(row=data_row, column=5).number_format


def test_table_creates_listobject_with_autofilter():
    data = build_excel_report(_minimal_spec())
    wb = load_workbook(BytesIO(data))
    ws = wb["Datos"]
    assert len(ws.tables) >= 1, "debe existir al menos un ListObject"
    table = next(iter(ws.tables.values()))
    assert table.autoFilter is not None


def test_freeze_panes_applied_when_requested():
    data = build_excel_report(_minimal_spec())
    wb = load_workbook(BytesIO(data))
    ws = wb["Datos"]
    assert ws.freeze_panes is not None


def test_header_uses_brand_orange():
    """Sin tabla (que aplica su propio estilo), el header debe llevar naranja Plataforma IA."""
    spec = ExcelReportSpec(
        filename="plain.xlsx",
        sheets=[SheetSpec(
            name="Datos",
            columns=[ColumnSpec(header="Col", key="c", type="text")],
            rows=[{"c": "x"}],
            table=False,
        )],
    )
    wb = load_workbook(BytesIO(build_excel_report(spec)))
    ws = wb["Datos"]
    fill_color = (ws.cell(row=1, column=1).fill.fgColor.rgb or "").upper()
    assert BRAND_ORANGE.lstrip("#").upper() in fill_color


def test_sanitize_applied_to_cell_values_during_build():
    spec = ExcelReportSpec(
        filename="injection.xlsx",
        sheets=[SheetSpec(
            name="Datos",
            columns=[ColumnSpec(header="Nota", key="nota", type="text")],
            rows=[{"nota": "=cmd|'/c calc'!A1"}],
        )],
    )
    wb = load_workbook(BytesIO(build_excel_report(spec)))
    ws = wb["Datos"]
    values = [c.value for row in ws.iter_rows(values_only=False) for c in row]
    sanitized_found = any(
        isinstance(v, str) and v.startswith("'=") for v in values
    )
    raw_formula = any(
        isinstance(v, str) and v.startswith("=cmd|") for v in values
    )
    assert sanitized_found or not raw_formula


def test_metadata_sets_author_and_title():
    data = build_excel_report(_minimal_spec())
    wb = load_workbook(BytesIO(data))
    assert wb.properties.creator == "Jorge Tapia"
    assert wb.properties.title


# ── inference ──────────────────────────────────────────────────────────────

def test_infer_spec_detects_numeric_columns():
    rows = [
        {"nombre": "a", "monto": 100, "porcentaje": 0.5},
        {"nombre": "b", "monto": 200, "porcentaje": 0.7},
    ]
    spec = infer_spec_from_rows(rows, filename="inferido.xlsx")
    types = {c.key: c.type for c in spec.sheets[0].columns}
    assert types["nombre"] == "text"
    assert types["monto"] == "number"


def test_infer_spec_handles_empty_rows():
    spec = infer_spec_from_rows([], filename="vacio.xlsx")
    assert spec.sheets[0].columns
    data = build_excel_report(spec)
    assert len(data) > 0


# ── schema validation ─────────────────────────────────────────────────────

def test_filename_must_end_with_xlsx():
    with pytest.raises(Exception):
        ExcelReportSpec(
            filename="no_extension",
            sheets=[SheetSpec(
                name="A",
                columns=[ColumnSpec(header="h", key="k")],
                rows=[],
            )],
        )


def test_sheet_name_max_length_31():
    long_name = "x" * 32
    with pytest.raises(Exception):
        SheetSpec(
            name=long_name,
            columns=[ColumnSpec(header="h", key="k")],
            rows=[],
        )


# ── nuevo engine: polars.write_excel + decisión inteligente ─────────────────

def _multi_metric_spec(report_type: str = "auto", **extras) -> ExcelReportSpec:
    return ExcelReportSpec(
        filename="dash.xlsx",
        sheets=[SheetSpec(
            name="Detalle",
            title="Producción Q1 2026",
            columns=[
                ColumnSpec(header="División", key="div", type="text"),
                ColumnSpec(header="Atribuciones", key="atr", type="text"),
                ColumnSpec(header="Toneladas", key="ton", type="integer"),
                ColumnSpec(header="USD MM", key="usd", type="currency"),
            ],
            rows=[
                {"div": "Norte", "atr": "Vicepresidencia Norte", "ton": 82500, "usd": 612.4},
                {"div": "RT", "atr": "Vicepresidencia Sur", "ton": 67800, "usd": 501.2},
                {"div": "Andina", "atr": "Vicepresidencia Centro", "ton": 54300, "usd": 412.0},
            ],
        )],
        report_type=report_type,
        author="Jorge Tapia",
        division="Norte",
        **extras,
    )


def test_polars_engine_renders_table_with_style():
    """polars.write_excel inserta un ListObject con estilo aplicado."""
    data = build_excel_report(_multi_metric_spec(report_type="data_table"))
    wb = load_workbook(BytesIO(data))
    ws = wb["Detalle"]
    assert len(ws.tables) >= 1
    table = next(iter(ws.tables.values()))
    assert table.tableStyleInfo is not None
    style = (table.tableStyleInfo.name or "").replace(" ", "")
    assert "TableStyle" in style


def test_autofit_widths_fit_long_strings():
    """autofit=True debe dimensionar 'Vicepresidencia Norte' sin truncar."""
    data = build_excel_report(_multi_metric_spec(report_type="data_table"))
    wb = load_workbook(BytesIO(data))
    ws = wb["Detalle"]
    # Columna B = "Atribuciones" (20+ chars en valor más largo).
    width_b = ws.column_dimensions.get("B")
    assert width_b is not None
    assert (width_b.width or 0) >= 14, (
        f"autofit insuficiente: ancho={width_b.width} para 'Vicepresidencia Norte'"
    )


def test_column_totals_sum_numeric_when_auto():
    """totals='auto' inserta una fila de totales SUM en numéricas."""
    data = build_excel_report(_multi_metric_spec(report_type="data_table"))
    wb = load_workbook(BytesIO(data))
    ws = wb["Detalle"]
    table = next(iter(ws.tables.values()))
    assert table.totalsRowShown or table.totalsRowCount, (
        "se esperaba totalsRow activo en la tabla"
    )


def test_conditional_format_applied_when_auto():
    """auto debe inferir un dict no vacío para columnas con varianza."""
    from app.agents.excel.auto_style import infer_conditional_formats
    spec = _multi_metric_spec(report_type="data_table")
    sheet = spec.sheets[0]
    inferred = infer_conditional_formats(sheet.rows, sheet.columns)
    assert inferred, "se esperaba al menos una regla inferida"
    # Y debe ser uno de los presets soportados por polars.
    assert all(v in {"data_bar", "3_color_scale", "2_color_scale"} for v in inferred.values())


def test_chart_embedded_when_spec_passed():
    """sheet_spec.chart agrega un chart al .xlsx (verificable en el zip)."""
    import zipfile
    spec = _multi_metric_spec(report_type="data_table")
    spec.sheets[0].chart = ChartSpec(
        type="bar",
        title="Producción por división",
        x_column="div",
        y_columns=["ton"],
    )
    data = build_excel_report(spec)
    with zipfile.ZipFile(BytesIO(data)) as zf:
        chart_files = [n for n in zf.namelist() if n.startswith("xl/charts/")]
    assert chart_files, "se esperaba al menos un xl/charts/chart*.xml en el zip"


def test_cover_sheet_present_with_metadata():
    """cover_sheet=True crea hoja Portada con título grande."""
    spec = _multi_metric_spec(report_type="kpi_dashboard", cover_sheet=True)
    data = build_excel_report(spec)
    wb = load_workbook(BytesIO(data))
    assert "Portada" in wb.sheetnames
    portada = wb["Portada"]
    cells_text = " ".join(
        str(c.value) for row in portada.iter_rows() for c in row if c.value
    )
    assert "Jorge Tapia" in cells_text or "Norte" in cells_text


def test_summary_sheet_with_kpi_tiles():
    """summary_sheet con KPIs explícitos crea hoja Resumen con labels visibles."""
    spec = _multi_metric_spec(report_type="kpi_dashboard")
    spec.summary_sheet = SummarySheetSpec(
        name="Resumen",
        kpis=[
            KPISpec(label="Producción total", value=204600, format="integer", color="green"),
            KPISpec(label="USD MM", value=1525.6, format="currency", delta=12.5, color="blue"),
        ],
    )
    data = build_excel_report(spec)
    wb = load_workbook(BytesIO(data))
    assert "Resumen" in wb.sheetnames
    resumen = wb["Resumen"]
    cells = " ".join(
        str(c.value) for row in resumen.iter_rows() for c in row if c.value
    )
    assert "Producción total" in cells
    assert "USD MM" in cells


def test_preset_executive_uses_navy_header():
    """style_preset='executive' aplica navy Plataforma IA en el header (vía table style)."""
    spec = _multi_metric_spec(report_type="kpi_dashboard")
    spec.style_preset = "executive"
    data = build_excel_report(spec)
    wb = load_workbook(BytesIO(data))
    ws = wb["Detalle"]
    table = next(iter(ws.tables.values()))
    # El estilo de tabla controla el color, así que validamos el nombre del estilo.
    assert "Medium7" in (table.tableStyleInfo.name or "").replace(" ", "")


def test_preset_modern_orange_default():
    spec = _multi_metric_spec(report_type="data_table")
    data = build_excel_report(spec)
    wb = load_workbook(BytesIO(data))
    ws = wb["Detalle"]
    table = next(iter(ws.tables.values()))
    assert "Medium" in (table.tableStyleInfo.name or "")


def test_report_type_auto_infers_ranking_for_small_tables():
    """1 categórica + 1 numérica + ≤30 filas → ranking."""
    columns = [
        ColumnSpec(header="Agente", key="agente", type="text"),
        ColumnSpec(header="Usos", key="usos", type="integer"),
    ]
    rows = [{"agente": f"a{i}", "usos": 100 - i * 5} for i in range(10)]
    inferred = infer_report_type(rows, columns)
    assert inferred == "ranking"


def test_report_type_auto_infers_timeseries_with_dates():
    """date + numeric → timeseries."""
    columns = [
        ColumnSpec(header="Fecha", key="fecha", type="date"),
        ColumnSpec(header="Producción", key="prod", type="number"),
    ]
    rows = [{"fecha": f"2026-{m:02d}-01", "prod": 100 + m * 10} for m in range(1, 7)]
    inferred = infer_report_type(rows, columns)
    assert inferred == "timeseries"


def test_chart_inferred_for_ranking_uses_bar():
    columns = [
        ColumnSpec(header="Agente", key="agente", type="text"),
        ColumnSpec(header="Usos", key="usos", type="integer"),
    ]
    rows = [{"agente": f"a{i}", "usos": 100 - i * 5} for i in range(8)]
    chart = infer_chart(rows, columns, "ranking")
    assert chart is not None
    assert chart.type == "bar"
    assert chart.x_column == "agente"
    assert chart.y_columns == ["usos"]


def test_kpis_inferred_from_numeric_columns():
    columns = [
        ColumnSpec(header="División", key="div", type="text"),
        ColumnSpec(header="Toneladas", key="ton", type="integer"),
        ColumnSpec(header="Margen", key="margen", type="percent"),
    ]
    rows = [
        {"div": "Norte", "ton": 82500, "margen": 0.32},
        {"div": "RT", "ton": 67800, "margen": 0.28},
    ]
    kpis = infer_kpis_from_rows(rows, columns)
    labels = [k.label for k in kpis]
    assert any("Toneladas" in l for l in labels)
    assert any("Margen" in l for l in labels)


def test_sparklines_render_in_sheet():
    """sparklines explícitas se insertan en una columna NUEVA (polars la inyecta)."""
    spec = ExcelReportSpec(
        filename="spark.xlsx",
        sheets=[SheetSpec(
            name="Mensual",
            columns=[
                ColumnSpec(header="División", key="div", type="text"),
                ColumnSpec(header="Ene", key="ene", type="integer"),
                ColumnSpec(header="Feb", key="feb", type="integer"),
                ColumnSpec(header="Mar", key="mar", type="integer"),
            ],
            rows=[
                {"div": "Norte", "ene": 80, "feb": 85, "mar": 90},
                {"div": "RT", "ene": 70, "feb": 65, "mar": 75},
            ],
            sparklines={"Tendencia": ["Ene", "Feb", "Mar"]},
            totals=None,
            conditional_formats=None,
        )],
        report_type="data_table",
    )
    data = build_excel_report(spec)
    assert len(data) > 0
    # openpyxl no parsea sparklines, pero el header "Tendencia" debe aparecer
    # como columna inyectada por polars.
    wb = load_workbook(BytesIO(data))
    ws = wb["Mensual"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "Tendencia" in headers


def test_backwards_compat_save_to_excel_inferred_spec():
    """infer_spec_from_rows + build_excel_report sigue funcionando con dicts simples."""
    rows = [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]
    spec = infer_spec_from_rows(rows, filename="legacy.xlsx")
    data = build_excel_report(spec)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames
    ws = wb[wb.sheetnames[0]]
    cells = [c.value for row in ws.iter_rows() for c in row]
    assert 1 in cells and 2 in cells


def test_executive_preset_navy_constant_present():
    """Smoke: el constante BRAND_NAVY se usa en el preset executive."""
    from app.agents.excel.theme import STYLE_PRESETS
    assert STYLE_PRESETS["executive"]["title_bg"] == BRAND_NAVY
